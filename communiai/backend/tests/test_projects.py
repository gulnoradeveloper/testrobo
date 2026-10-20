import io
import zipfile

from conftest import create_project, register


def _docx_bytes(text):
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pptx_bytes(text):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "SmartQueue"
    slide.placeholders[1].text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pdf_bytes(text):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def test_project_crud(client):
    h, _ = register(client)
    p = create_project(client, h)
    assert p["title"].startswith("SmartQueue")
    r = client.patch(f"/api/v1/projects/{p['id']}", json={"features": "Navbat olish, SMS xabarnoma"}, headers=h)
    assert r.status_code == 200 and r.json()["features"] == "Navbat olish, SMS xabarnoma"
    lst = client.get("/api/v1/projects", headers=h).json()
    assert lst["total"] == 1 and lst["page"] == 1
    assert client.delete(f"/api/v1/projects/{p['id']}", headers=h).status_code == 204
    assert client.get(f"/api/v1/projects/{p['id']}", headers=h).status_code == 404
    assert client.get("/api/v1/projects", headers=h).json()["total"] == 0


def test_project_required_fields(client):
    h, _ = register(client)
    r = client.post("/api/v1/projects", json={"title": "AB", "category": "web"}, headers=h)
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "VALIDATION_ERROR"
    r = client.post("/api/v1/projects", json={"title": "Loyiha", "category": "nomalum", "description": "x" * 30,
                                              "problem": "muammo bor juda katta"}, headers=h)
    assert r.status_code == 422


def test_upload_and_extract_all_formats(client):
    h, _ = register(client)
    p = create_project(client, h)
    text = "SmartQueue tizimi navbat raqamini atomar tranzaksiya bilan beradi"
    files = {
        "a.txt": (text.encode(), "text/plain"),
        "b.docx": (_docx_bytes(text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "c.pptx": (_pptx_bytes(text), "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
        "d.pdf": (_pdf_bytes("SmartQueue navbat tizimi hujjati PDF formatida"), "application/pdf"),
    }
    for name, (data, mime) in files.items():
        r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": (name, data, mime)}, headers=h)
        assert r.status_code == 201, (name, r.text)
        body = r.json()
        assert body["extraction_status"] == "completed", (name, body)
        assert "SmartQueue" in body["extracted_preview"]
    assert len(client.get(f"/api/v1/projects/{p['id']}/files", headers=h).json()["items"]) == 4


def test_duplicate_upload_is_idempotent(client):
    h, _ = register(client)
    p = create_project(client, h)
    data = b"Bir xil fayl mazmuni, ikki marta yuborildi. " * 3
    a = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("x.txt", data, "text/plain")}, headers=h)
    b = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("x.txt", data, "text/plain")}, headers=h)
    assert a.json()["id"] == b.json()["id"]


def test_reject_wrong_format_and_spoofed_signature(client):
    h, _ = register(client)
    p = create_project(client, h)
    r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("virus.exe", b"MZ\x90\x00", "application/octet-stream")}, headers=h)
    assert r.status_code == 415 and r.json()["error"]["code"] == "UNSUPPORTED_TYPE"
    r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("fake.pdf", b"not a pdf at all", "application/pdf")}, headers=h)
    assert r.status_code == 400 and r.json()["error"]["code"] == "SIGNATURE_MISMATCH"
    r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("x.docx", b"PK\x03\x04garbage", "application/zip")}, headers=h)
    assert r.status_code == 400
    r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("x.pdf", b"%PDF-1.4", "image/png")}, headers=h)
    assert r.status_code == 415 and r.json()["error"]["code"] == "MIME_MISMATCH"


def test_reject_too_large(client, monkeypatch):
    from app.core.config import get_settings
    monkeypatch.setattr(get_settings(), "MAX_UPLOAD_SIZE_MB", 1)
    h, _ = register(client)
    p = create_project(client, h)
    big = b"a" * (1024 * 1024 + 10)
    r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("big.txt", big, "text/plain")}, headers=h)
    assert r.status_code == 413 and r.json()["error"]["code"] == "FILE_TOO_LARGE"


def test_reject_zip_bomb(client):
    h, _ = register(client)
    p = create_project(client, h)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml", b"0" * (200 * 1024 * 1024))
    r = client.post(f"/api/v1/projects/{p['id']}/files",
                    files={"file": ("bomb.docx", buf.getvalue(), "application/zip")}, headers=h)
    assert r.status_code == 400 and r.json()["error"]["code"] == "ARCHIVE_LIMIT"


def test_empty_text_pdf_reports_no_text(client):
    h, _ = register(client)
    p = create_project(client, h)
    r = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("scan.pdf", _pdf_bytes(""), "application/pdf")},
                    headers=h)
    assert r.status_code == 201
    assert r.json()["extraction_status"] == "empty"
    assert "OCR" in r.json()["extraction_error"]


def test_delete_file(client):
    h, _ = register(client)
    p = create_project(client, h)
    f = client.post(f"/api/v1/projects/{p['id']}/files", files={"file": ("a.txt", b"Matn " * 10, "text/plain")},
                    headers=h).json()
    assert client.delete(f"/api/v1/projects/{p['id']}/files/{f['id']}", headers=h).status_code == 204
    assert client.get(f"/api/v1/projects/{p['id']}/files", headers=h).json()["items"] == []


def test_filename_is_sanitized(client):
    h, _ = register(client)
    p = create_project(client, h)
    r = client.post(f"/api/v1/projects/{p['id']}/files",
                    files={"file": ("../../etc/<script>.txt", b"Matn mazmuni yetarli uzun", "text/plain")}, headers=h)
    assert r.status_code == 201
    assert "/" not in r.json()["original_filename"] and "<" not in r.json()["original_filename"]
