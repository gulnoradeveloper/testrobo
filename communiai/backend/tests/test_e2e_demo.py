"""End-to-end run of the acceptance demo scenario (TZ section 24) through the HTTP API.

The AI provider is the test-only scripted provider; everything else (HTTP, auth, DB, storage,
extraction, prompts, validation, scoring, jobs) is the real stack.
"""
import io

from conftest import PASSWORD


def _pptx():
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "SmartQueue"
    s.placeholders[1].text = "Navbat raqami PostgreSQL ketma-ketligi orqali atomar beriladi. Talabalar SMS oladi."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_demo_scenario(client, ai):
    # 1. register
    r = client.post("/api/v1/auth/register", json={
        "first_name": "Dilnoza", "last_name": "Rahimova", "email": "dilnoza@example.uz", "password": PASSWORD,
        "password_confirm": PASSWORD, "university": "TATU", "study_program": "60610100"})
    assert r.status_code == 201
    h = {"Authorization": f"Bearer {r.json()['access_token']}"}
    client.patch("/api/v1/me", json={"ai_consent": True}, headers=h)

    # 2. create SmartQueue project
    p = client.post("/api/v1/projects", headers=h, json={
        "title": "SmartQueue — elektron navbat tizimi", "category": "info_system",
        "description": "Talabalar uchun dekanatga onlayn navbat olish va navbat holatini kuzatish tizimi.",
        "problem": "Dekanat oldida uzun navbatlar va talabalar vaqtining yo'qotilishi.",
        "technologies": "Flutter, FastAPI, PostgreSQL"}).json()

    # 3. upload a PPTX
    f = client.post(f"/api/v1/projects/{p['id']}/files", headers=h, files={"file": (
        "smartqueue.pptx", _pptx(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}).json()
    assert f["extraction_status"] == "completed" and "atomar" in f["extracted_preview"]

    # 4-5. AI expertise: strengths, weaknesses, recommendations, expected questions
    a = client.post(f"/api/v1/projects/{p['id']}/analyze", headers=h).json()
    a = client.get(f"/api/v1/analyses/{a['id']}", headers=h).json()
    assert a["status"] == "completed"
    res = a["result"]
    assert res["strengths"] and res["weaknesses"] and res["project_recommendations"]
    assert len(res["expected_questions"]) >= 5
    assert "atomar beriladi" in ai.calls[-1]["user"]  # the PPTX content reached the analysis

    # 6-7. hackathon mode, 5 dynamic questions
    d = client.post("/api/v1/defenses", headers=h,
                    json={"project_id": p["id"], "mode": "hackathon", "question_limit": 5}).json()
    client.post(f"/api/v1/defenses/{d['id']}/start", headers=h)
    # 8-9. answer each question and get short feedback
    for i in range(5):
        cur = client.get(f"/api/v1/defenses/{d['id']}/current-question", headers=h).json()
        if i:
            assert cur["last_feedback"]["feedback"]["strength"]
        client.post(f"/api/v1/defenses/{d['id']}/answers", headers=h, json={
            "question_id": cur["question"]["id"],
            "answer_text": f"{i + 1}-javob: navbat raqami ketma-ketlik va unique cheklov bilan beriladi."})
    # 10. final communicative assessment
    assert client.post(f"/api/v1/defenses/{d['id']}/complete", headers=h).json()["status"] == "EVALUATED"
    result = client.get(f"/api/v1/defenses/{d['id']}/results", headers=h).json()
    assert result["result"]["level"] and result["plan_id"]

    # 11-13. individual practice, revised answer, before/after comparison
    plan = client.get(f"/api/v1/practice-plans/{result['plan_id']}", headers=h).json()
    client.post(f"/api/v1/practice-plans/{plan['id']}/attempts", headers=h, json={
        "exercise_index": 0, "revised_answer": "Avval savolni qayta ifodalayman: takroriy raqamlarning oldini olish. "
                                               "Javob: ketma-ketlik, tranzaksiya va unique indeks."})
    prog = client.get(f"/api/v1/practice-plans/{plan['id']}/progress", headers=h).json()
    assert prog["rows"][0]["before"] is not None and prog["rows"][0]["after"] is not None

    # 14-15. log out, log in again: everything is still there
    client.post("/api/v1/auth/logout", headers=h, json={"refresh_token": r.json()["refresh_token"]})
    tok = client.post("/api/v1/auth/login", json={"email": "dilnoza@example.uz", "password": PASSWORD}).json()
    h2 = {"Authorization": f"Bearer {tok['access_token']}"}
    dash = client.get("/api/v1/me/dashboard", headers=h2).json()
    assert dash["projects_count"] == 1 and dash["completed_defenses"] == 1 and dash["latest_result"]
    assert client.get("/api/v1/me/progress", headers=h2).json()["practice"]
    hist = {i["action"] for i in client.get("/api/v1/me/history", headers=h2).json()["items"]}
    assert {"project.create", "file.upload", "analysis.request", "defense.complete", "practice.attempt"} <= hist
