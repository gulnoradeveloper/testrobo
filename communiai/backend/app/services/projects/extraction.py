"""Upload validation and text extraction for PDF, DOCX, PPTX and TXT.

Files are only parsed, never executed. Office files are ZIP containers, so archive limits
(entry count, total uncompressed size, compression ratio) are enforced before parsing.
"""
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass

from app.core.config import get_settings

ALLOWED = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".txt": "text/plain",
}
# Declared MIME types accepted from clients for each extension (some clients send generic types).
ACCEPTED_DECLARED = {
    ".pdf": {"application/pdf", "application/octet-stream", "application/x-pdf"},
    ".docx": {ALLOWED[".docx"], "application/octet-stream", "application/zip"},
    ".pptx": {ALLOWED[".pptx"], "application/octet-stream", "application/zip"},
    ".txt": {"text/plain", "application/octet-stream"},
}


class UploadRejected(Exception):
    def __init__(self, code: str, message: str):
        self.code, self.message = code, message
        super().__init__(message)


@dataclass
class ExtractionResult:
    status: str  # completed | empty | failed
    text: str
    error: str | None = None


def safe_filename(name: str) -> str:
    name = unicodedata.normalize("NFKC", name or "fayl")
    name = name.replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^\w.\- ()'ʻ‘]+", "_", name, flags=re.UNICODE).strip(" .")
    return (name or "fayl")[:200]


def extension_of(name: str) -> str:
    m = re.search(r"(\.[A-Za-z0-9]+)$", name or "")
    return m.group(1).lower() if m else ""


def _check_zip(data: bytes, required_prefix: str) -> None:
    s = get_settings()
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise UploadRejected("INVALID_FILE", "Fayl buzilgan yoki formati noto'g'ri") from e
    infos = zf.infolist()
    if len(infos) > s.MAX_ARCHIVE_ENTRIES:
        raise UploadRejected("ARCHIVE_LIMIT", "Fayl ichidagi elementlar soni juda ko'p")
    total = sum(i.file_size for i in infos)
    if total > s.MAX_ARCHIVE_UNCOMPRESSED_MB * 1024 * 1024:
        raise UploadRejected("ARCHIVE_LIMIT", "Fayl ochilganda hajmi ruxsat etilgan chegaradan oshadi")
    for i in infos:
        if i.compress_size and i.file_size / max(i.compress_size, 1) > 200 and i.file_size > 5 * 1024 * 1024:
            raise UploadRejected("ARCHIVE_LIMIT", "Faylda shubhali siqilish nisbati aniqlandi")
    names = {i.filename for i in infos}
    if "[Content_Types].xml" not in names or not any(n.startswith(required_prefix) for n in names):
        raise UploadRejected("INVALID_FILE", "Fayl mazmuni kengaytmasiga mos emas")


def validate_upload(filename: str, declared_mime: str | None, data: bytes) -> tuple[str, str]:
    """Return (extension, canonical_mime) or raise UploadRejected."""
    s = get_settings()
    if not data:
        raise UploadRejected("EMPTY_FILE", "Fayl bo'sh")
    if len(data) > s.max_upload_bytes:
        raise UploadRejected("FILE_TOO_LARGE", f"Fayl hajmi {s.MAX_UPLOAD_SIZE_MB} MB dan oshmasligi kerak")
    ext = extension_of(filename)
    if ext not in ALLOWED:
        raise UploadRejected("UNSUPPORTED_TYPE", "Faqat PDF, DOCX, PPTX va TXT fayllar qabul qilinadi")
    if declared_mime and declared_mime.split(";")[0].strip().lower() not in ACCEPTED_DECLARED[ext]:
        raise UploadRejected("MIME_MISMATCH", "Fayl turi (MIME) kengaytmasiga mos emas")
    # Signature (magic bytes) checks
    if ext == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF-"):
            raise UploadRejected("SIGNATURE_MISMATCH", "Fayl PDF formatida emas")
    elif ext in (".docx", ".pptx"):
        if not data.startswith(b"PK\x03\x04"):
            raise UploadRejected("SIGNATURE_MISMATCH", "Fayl Office (ZIP) formatida emas")
        _check_zip(data, "word/" if ext == ".docx" else "ppt/")
    elif ext == ".txt":
        if b"\x00" in data[:8192]:
            raise UploadRejected("SIGNATURE_MISMATCH", "Matnli fayl ikkilik ma'lumot o'z ichiga oladi")
        if data.startswith((b"MZ", b"\x7fELF", b"PK\x03\x04", b"%PDF")):
            raise UploadRejected("SIGNATURE_MISMATCH", "Fayl matnli fayl emas")
    return ext, ALLOWED[ext]


def _clean(text: str) -> str:
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise UploadRejected("ENCRYPTED", "Parol bilan himoyalangan PDF qo'llab-quvvatlanmaydi") from e
    out = []
    for i, page in enumerate(reader.pages[:300]):
        out.append(f"[Sahifa {i + 1}]\n" + (page.extract_text() or ""))
    return "\n".join(out)


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        texts.append(" | ".join(cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"(Izoh: {notes})")
        out.append(f"[Slayd {i}]\n" + "\n".join(texts))
    return "\n".join(out)


def _txt(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1251", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_text(ext: str, data: bytes) -> ExtractionResult:
    s = get_settings()
    try:
        raw = {".pdf": _pdf, ".docx": _docx, ".pptx": _pptx, ".txt": _txt}[ext](data)
    except UploadRejected as e:
        return ExtractionResult("failed", "", e.message)
    except Exception:
        return ExtractionResult("failed", "", "Fayldan matn ajratib bo'lmadi: fayl buzilgan yoki qo'llab-quvvatlanmaydi")
    text = _clean(raw)
    meaningful = re.sub(r"\[(Sahifa|Slayd) \d+\]", "", text).strip()
    if len(meaningful) < 20:
        return ExtractionResult(
            "empty", "", "Fayldan matn topilmadi. Skanerlangan PDF yoki rasm bo'lishi mumkin (OCR MVP da mavjud emas).")
    if len(text) > s.MAX_EXTRACTED_CHARS:
        text = text[: s.MAX_EXTRACTED_CHARS]
    return ExtractionResult("completed", text)
